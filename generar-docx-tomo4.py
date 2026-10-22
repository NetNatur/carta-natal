# Genera el DOCX 6x9 del tomo IV a partir de volumen4 - prospectiva.md,
# usando el tomo I maquetado como plantilla (página, márgenes y estilos),
# y añade la paginación en el pie (portada sin número).
# Uso: python generar-docx-tomo4.py
import re
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

PLANTILLA = r'projects\agente-astrologia\registro\volumen1 - final 6x9 - con indice.docx'
MD = r'projects\agente-astrologia\registro\volumen4 - prospectiva.md'
SALIDA = r'projects\agente-astrologia\registro\volumen4 - final 6x9.docx'

doc = docx.Document(PLANTILLA)

# 1) vaciar el cuerpo (conserva sectPr, estilos y configuración de página)
for p in list(doc.paragraphs):
    p._p.getparent().remove(p._p)

# 2) volcar el md
lineas = open(MD, encoding='utf8').read().split('\n')

def agregar_parrafo(texto, estilo=None, justificado=True):
    p = doc.add_paragraph()
    if estilo:
        p.style = doc.styles[estilo]
    if texto:
        # parseo de **negritas**
        partes = re.split(r'\*\*', texto)
        for i, chunk in enumerate(partes):
            if not chunk:
                continue
            run = p.add_run(chunk)
            run.bold = (i % 2 == 1)
            if not estilo:
                run.font.name = 'Georgia'
                run.font.size = Pt(11)
    if not estilo and justificado:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

primera_linea = True
for linea in lineas:
    t = linea.rstrip()
    if primera_linea and t.startswith('# '):
        agregar_parrafo(t[2:].strip(), 'Title')
        primera_linea = False
        continue
    if t.startswith('Tomo IV'):
        agregar_parrafo(t, 'Subtitle')
        continue
    if t.startswith('### '):
        agregar_parrafo(t[4:].strip(), 'Heading 3')
        continue
    if t.startswith('## '):
        agregar_parrafo(t[3:].strip(), 'Heading 2')
        continue
    if t.startswith('# '):
        agregar_parrafo(t[2:].strip(), 'Heading 1')
        continue
    if t == '':
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        continue
    if t.startswith('•'):
        agregar_parrafo(t, justificado=False).paragraph_format.left_indent = Inches(0.25)
        continue
    agregar_parrafo(t)

# 2b) el 'ÍNDICE' del md se escribió sin '#': convertirlo en Heading 1 (ancla del índice)
for p in doc.paragraphs:
    if p.text.strip() == 'ÍNDICE' and p.style.name != 'Heading 1':
        p.style = doc.styles['Heading 1']
        break

# 3) paginación: pie centrado con campo PAGE, portada sin número
sec = doc.sections[0]
sec.different_first_page_header_footer = True
pie = sec.footer
if pie.is_linked_to_previous:
    pie.is_linked_to_previous = False
parr = pie.paragraphs[0] if pie.paragraphs else pie.add_paragraph()
parr.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = parr.add_run()
f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'PAGE'
f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
run._r.append(f1); run._r.append(it); run._r.append(f2)
run.font.name = 'Georgia'; run.font.size = Pt(10)
run.font.color.rgb = RGBColor.from_string('535356')
if sec.first_page_footer.is_linked_to_previous:
    sec.first_page_footer.is_linked_to_previous = False

doc.save(SALIDA)
print('OK:', SALIDA)

# verificación
d2 = docx.Document(SALIDA)
from collections import Counter
c = Counter(p.style.name for p in d2.paragraphs)
print('parrafos:', len(d2.paragraphs), '| estilos:', dict(c))
s = d2.sections[0]
print('pagina:', round(s.page_width.mm, 1), 'x', round(s.page_height.mm, 1), 'mm')