# Índice estático con páginas reales para el tomo IV (adaptado del tomo I).
# Exporta el PDF con LibreOffice, detecta la página de cada título, escribe el
# índice con tabulaciones de puntos e itera hasta converger.
# Uso: python indice-tomo4.py
import re, subprocess, os, shutil
import docx
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor
from pypdf import PdfReader

REGISTRO = r'projects\agente-astrologia\registro'
BASE = os.path.join(REGISTRO, 'volumen4 - final 6x9.docx')
DOCX = os.path.join(REGISTRO, 'volumen4 - final 6x9 - con indice.docx')
PDF  = os.path.join(REGISTRO, 'volumen4 - final 6x9 - con indice.pdf')
SOFFICE = r'C:\Program Files\LibreOffice\program\soffice.exe'
PRIMER_TITULO = 'PARTE I · EL MÉTODO DE LA PROSPECTIVA'

shutil.copyfile(BASE, DOCX)

def exportar_pdf():
    subprocess.run([SOFFICE, '--headless', '--convert-to', 'pdf',
                    '--outdir', REGISTRO, DOCX], check=True,
                   capture_output=True, timeout=900)

def norm(s):
    return re.sub(r'[^\w]', '', s, flags=re.UNICODE).lower()

def titulos_docx():
    doc = docx.Document(DOCX)
    titulos = []
    for p in doc.paragraphs:
        t = p.text.strip()
        st = p.style.name
        if not t:
            continue
        if st == 'Heading 1' and t != 'ÍNDICE':
            titulos.append((1, t))
        elif st == 'Heading 2':
            titulos.append((2, t))
        elif st == 'Heading 3':
            titulos.append((3, t))
    return titulos

def paginas_de_titulos(titulos):
    r = PdfReader(PDF)
    crudos = [pg.extract_text() or '' for pg in r.pages]
    n_pags = len(crudos)
    lineas = [[norm(l) for l in p.splitlines() if l.strip()] for p in crudos]
    frag = [norm(p) for p in crudos]
    t0 = norm(PRIMER_TITULO)
    inicio = 0
    for i in range(n_pags):
        if any(ln == t0 for ln in lineas[i]):
            inicio = i
            break
    paginas = []
    anterior = inicio
    for nivel, t in titulos:
        nt = norm(t)
        pag = None
        for i in range(max(anterior - 1, 0), n_pags):
            if any(ln == nt for ln in lineas[i]):
                pag = i + 1
                break
        if pag is None:
            for i in range(max(anterior - 1, 0), n_pags):
                if nt in frag[i]:
                    pag = i + 1
                    break
        if pag is None:
            pag = anterior + 1
        paginas.append(pag)
        anterior = pag
    return paginas

def escribir_indice(titulos, paginas):
    doc = docx.Document(DOCX)
    h1 = None
    h1_parte = None
    toc_p = None
    for p in doc.paragraphs:
        t = p.text.strip()
        if t == 'ÍNDICE' and p.style.name == 'Heading 1':
            h1 = p
        elif norm(t) == norm(PRIMER_TITULO) and p.style.name == 'Heading 1':
            h1_parte = p
        if 'fldSimple' in p._p.xml:
            toc_p = p
    if h1 is None or h1_parte is None:
        raise RuntimeError('no encuentro las anclas del índice')
    if toc_p is not None:
        toc_p._p.getparent().remove(toc_p._p)
    node = h1._p.getnext()
    while node is not None and node is not h1_parte._p:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt
    sec = doc.sections[0]
    usable = sec.page_width - sec.left_margin - sec.right_margin
    nuevo = []
    for (nivel, t), pag in zip(titulos, paginas):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        if nivel == 1:
            pf.left_indent = Inches(0)
            run = p.add_run(t + '\t' + str(pag))
            run.bold = True
        else:
            pf.left_indent = Inches(0.22 if nivel == 2 else 0.44)
            run = p.add_run(t + '\t' + str(pag))
            run.bold = False
        run.font.name = 'Georgia'
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0, 0, 0)
        nuevo.append(p)
    for p in reversed(nuevo):
        h1._p.addnext(p._p)
    doc.save(DOCX)

titulos = titulos_docx()
print('títulos:', len(titulos))
exportar_pdf()
paginas = paginas_de_titulos(titulos)
print('pasada 1 (sin índice):', paginas[:4], '...', paginas[-3:])
for it in range(2, 6):
    escribir_indice(titulos, paginas)
    exportar_pdf()
    nuevas = paginas_de_titulos(titulos)
    dif = sum(1 for a, b in zip(paginas, nuevas) if a != b)
    print(f'pasada {it}: {dif} diferencias |', nuevas[:4], '...', nuevas[-3:])
    if dif == 0:
        break
    paginas = nuevas

import random
random.seed(41)
r = PdfReader(PDF)
muestras = random.sample(list(zip(titulos, paginas)), 6)
ok = 0
for (nivel, t), pag in muestras:
    nt = norm(t)
    lineas_pag = [norm(l) for l in (r.pages[pag - 1].extract_text() or '').splitlines()]
    if any(ln == nt for ln in lineas_pag):
        ok += 1
    else:
        print(f'  REVISAR [{pag}] {t}')
print(f'verificación: {ok}/6 | páginas totales: {len(r.pages)}')